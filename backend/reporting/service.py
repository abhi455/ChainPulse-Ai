from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any


@dataclass
class ReportSection:
    title: str
    text: str = ""
    table: list[dict[str, Any]] = field(
        default_factory=list
    )


@dataclass
class ReportDefinition:
    title: str
    subtitle: str = ""
    sections: list[ReportSection] = field(
        default_factory=list
    )


class ReportService:

    def __init__(
        self,
        output_dir: str = "storage/reports",
    ):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(
            parents=True,
            exist_ok=True,
        )

    def generate_docx(
        self,
        report: ReportDefinition,
        filename: str,
    ) -> Path:

        from docx import Document

        path = self.output_dir / filename

        document = Document()

        document.add_heading(
            report.title,
            0,
        )

        if report.subtitle:
            document.add_paragraph(
                report.subtitle
            )

        for section in report.sections:

            document.add_heading(
                section.title,
                level=1,
            )

            if section.text:
                document.add_paragraph(
                    section.text
                )

            if section.table:

                columns = list(
                    section.table[0].keys()
                )

                table = document.add_table(
                    rows=1,
                    cols=len(columns),
                )

                for i, column in enumerate(columns):
                    table.rows[0].cells[i].text = str(
                        column
                    )

                for row in section.table:
                    cells = table.add_row().cells

                    for i, column in enumerate(columns):
                        cells[i].text = str(
                            row.get(column, "")
                        )

        document.save(path)

        return path

    def generate_pdf(
        self,
        report: ReportDefinition,
        filename: str,
    ) -> Path:

        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import (
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )
        from reportlab.lib import colors

        path = self.output_dir / filename

        document = SimpleDocTemplate(
            str(path),
            pagesize=A4,
        )

        styles = getSampleStyleSheet()

        story = [
            Paragraph(
                report.title,
                styles["Title"],
            )
        ]

        if report.subtitle:
            story.append(
                Paragraph(
                    report.subtitle,
                    styles["Normal"],
                )
            )

        story.append(Spacer(1, 12))

        for section in report.sections:

            story.append(
                Paragraph(
                    section.title,
                    styles["Heading2"],
                )
            )

            if section.text:
                story.append(
                    Paragraph(
                        section.text,
                        styles["BodyText"],
                    )
                )

            if section.table:

                columns = list(
                    section.table[0].keys()
                )

                data = [
                    columns
                ]

                for row in section.table:
                    data.append(
                        [
                            str(row.get(column, ""))
                            for column in columns
                        ]
                    )

                table = Table(
                    data,
                    repeatRows=1,
                )

                table.setStyle(
                    TableStyle(
                        [
                            (
                                "BACKGROUND",
                                (0, 0),
                                (-1, 0),
                                colors.lightgrey,
                            ),
                            (
                                "GRID",
                                (0, 0),
                                (-1, -1),
                                0.5,
                                colors.grey,
                            ),
                            (
                                "VALIGN",
                                (0, 0),
                                (-1, -1),
                                "TOP",
                            ),
                        ]
                    )
                )

                story.append(table)
                story.append(
                    Spacer(1, 10)
                )

        document.build(story)

        return path

    def generate_xlsx(
        self,
        report: ReportDefinition,
        filename: str,
    ) -> Path:

        import pandas as pd

        path = self.output_dir / filename

        with pd.ExcelWriter(
            path,
            engine="xlsxwriter",
        ) as writer:

            summary = pd.DataFrame(
                [
                    {
                        "Report": report.title,
                        "Subtitle": report.subtitle,
                    }
                ]
            )

            summary.to_excel(
                writer,
                sheet_name="Summary",
                index=False,
            )

            for section in report.sections:

                if section.table:
                    df = pd.DataFrame(
                        section.table
                    )

                    safe_name = (
                        section.title[:31]
                    )

                    df.to_excel(
                        writer,
                        sheet_name=safe_name,
                        index=False,
                    )

        return path
