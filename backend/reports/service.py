from __future__ import annotations

from io import BytesIO
from datetime import datetime

import matplotlib.pyplot as plt
from docx import Document  # pyright: ignore[reportMissingImports]
from docx.shared import Inches  # pyright: ignore[reportMissingImports]
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors  # pyright: ignore[reportMissingModuleSource]
from reportlab.lib.pagesizes import A4  # pyright: ignore[reportMissingModuleSource]
from reportlab.lib.styles import getSampleStyleSheet  # pyright: ignore[reportMissingModuleSource]
from reportlab.platypus import (  # pyright: ignore[reportMissingModuleSource]
    Image,  # pyright: ignore[reportMissingModuleSource]
    Paragraph,  # pyright: ignore[reportMissingModuleSource]
    SimpleDocTemplate,  # pyright: ignore[reportMissingModuleSource]
    Spacer,  # pyright: ignore[reportMissingModuleSource]
    Table,  # pyright: ignore[reportMissingModuleSource]
    TableStyle,  # pyright: ignore[reportMissingModuleSource]
)

from backend.dashboard import DashboardService


class ReportService:
    """
    Unified ChainPulse AI reporting/export service.

    All export formats use the same dashboard summary so that
    Excel, PDF and Word remain synchronized.
    """

    # =========================================================
    # DATA
    # =========================================================

    @staticmethod
    def _summary(db, organization_id):
        return DashboardService(db).summary(
            organization_id
        )

    @staticmethod
    def _metrics(summary: dict) -> list[tuple[str, object]]:
        """
        Single source of truth for exported KPI values.
        """

        risk_level = summary.get(
            "overall_risk_level",
            "N/A",
        )

        risk_score = summary.get(
            "overall_risk_score",
            0,
        )

        if risk_level == "N/A":
            overall_risk = "N/A"
        else:
            overall_risk = (
                f"{float(risk_score):.2f} "
                f"{str(risk_level).upper()}"
            )

        supplier_reliability = summary.get(
            "average_supplier_reliability",
            0,
        )

        if summary.get("supplier_count", 0) == 0:
            supplier_reliability_display = "N/A"
        else:
            supplier_reliability_display = (
                f"{float(supplier_reliability):.2f}%"
            )

        return [
            (
                "Products",
                summary.get("product_count", 0),
            ),
            (
                "Active Products",
                summary.get("active_product_count", 0),
            ),
            (
                "Demand Records",
                summary.get("demand_record_count", 0),
            ),
            (
                "Total Demand",
                summary.get("total_demand_quantity", 0),
            ),
            (
                "Total Revenue",
                summary.get("total_revenue", 0),
            ),
            (
                "Inventory On Hand",
                summary.get("inventory_on_hand", 0),
            ),
            (
                "Inventory Reserved",
                summary.get("inventory_reserved", 0),
            ),
            (
                "Inventory Available",
                summary.get("inventory_available", 0),
            ),
            (
                "Inventory Below Reorder",
                summary.get("inventory_below_reorder", 0),
            ),
            (
                "Inventory Stockouts",
                summary.get("inventory_stockouts", 0),
            ),
            (
                "Suppliers",
                summary.get("supplier_count", 0),
            ),
            (
                "Supplier Reliability",
                supplier_reliability_display,
            ),
            (
                "Overall Risk",
                overall_risk,
            ),
            (
                "Demand Risk",
                summary.get(
                    "demand_risk_level",
                    "N/A",
                ),
            ),
            (
                "Inventory Risk",
                summary.get(
                    "inventory_risk_level",
                    "N/A",
                ),
            ),
            (
                "Supplier Risk",
                summary.get(
                    "supplier_risk_level",
                    "N/A",
                ),
            ),
            (
                "Forecast Count",
                summary.get(
                    "latest_forecast_count",
                    0,
                ),
            ),
            (
                "Scenarios",
                summary.get(
                    "scenario_count",
                    0,
                ),
            ),
            (
                "Simulation Runs",
                summary.get(
                    "simulation_run_count",
                    0,
                ),
            ),
        ]

    @staticmethod
    def _has_operational_data(summary: dict) -> bool:
        return any(
            [
                summary.get("product_count", 0),
                summary.get("demand_record_count", 0),
                summary.get("inventory_record_count", 0),
                summary.get("supplier_count", 0),
            ]
        )

    # =========================================================
    # CHART
    # =========================================================

    @staticmethod
    def _create_kpi_chart(summary: dict) -> BytesIO | None:
        """
        Create a simple KPI chart.

        Returns None when there is no operational data.
        """

        if not ReportService._has_operational_data(summary):
            return None

        labels = [
            "Products",
            "Demand",
            "Inventory",
            "Suppliers",
        ]

        values = [
            float(summary.get("product_count", 0)),
            float(summary.get("demand_record_count", 0)),
            float(summary.get("inventory_record_count", 0)),
            float(summary.get("supplier_count", 0)),
        ]

        figure = plt.figure(figsize=(8, 4.5))
        axis = figure.add_subplot(111)

        axis.bar(labels, values)

        axis.set_title(
            "ChainPulse Operational Overview"
        )
        axis.set_ylabel("Count")

        axis.grid(
            axis="y",
            alpha=0.25,
        )

        figure.tight_layout()

        buffer = BytesIO()

        figure.savefig(
            buffer,
            format="png",
            dpi=160,
            bbox_inches="tight",
        )

        plt.close(figure)

        buffer.seek(0)

        return buffer

    # =========================================================
    # PDF
    # =========================================================

    @staticmethod
    def generate_pdf(
        db,
        organization_id,
        request,
    ) -> BytesIO:

        summary = ReportService._summary(
            db,
            organization_id,
        )

        buffer = BytesIO()

        document = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            rightMargin=40,
            leftMargin=40,
            topMargin=40,
            bottomMargin=40,
        )

        styles = getSampleStyleSheet()

        story = []

        # -----------------------------------------------------
        # TITLE
        # -----------------------------------------------------

        story.append(
            Paragraph(
                request.title,
                styles["Title"],
            )
        )

        story.append(
            Paragraph(
                "Generated by ChainPulse AI",
                styles["Normal"],
            )
        )

        story.append(
            Paragraph(
                datetime.now().strftime(
                    "Generated: %d %B %Y, %H:%M"
                ),
                styles["Normal"],
            )
        )

        story.append(
            Spacer(1, 20)
        )

        # -----------------------------------------------------
        # DASHBOARD
        # -----------------------------------------------------

        if request.include_dashboard:

            story.append(
                Paragraph(
                    "Dashboard Overview",
                    styles["Heading2"],
                )
            )

            metrics = ReportService._metrics(
                summary
            )

            table_data = [
                ["Metric", "Value"]
            ]

            for label, value in metrics:
                table_data.append(
                    [
                        str(label),
                        str(value),
                    ]
                )

            table = Table(
                table_data,
                colWidths=[250, 200],
            )

            table.setStyle(
                TableStyle(
                    [
                        (
                            "BACKGROUND",
                            (0, 0),
                            (-1, 0),
                            colors.HexColor("#1F2937"),
                        ),
                        (
                            "TEXTCOLOR",
                            (0, 0),
                            (-1, 0),
                            colors.white,
                        ),
                        (
                            "FONTNAME",
                            (0, 0),
                            (-1, 0),
                            "Helvetica-Bold",
                        ),
                        (
                            "GRID",
                            (0, 0),
                            (-1, -1),
                            0.5,
                            colors.grey,
                        ),
                        (
                            "VALIGN",
                            (0, 0),
                            (-1, -1),
                            "MIDDLE",
                        ),
                        (
                            "PADDING",
                            (0, 0),
                            (-1, -1),
                            6,
                        ),
                    ]
                )
            )

            story.append(table)

            story.append(
                Spacer(1, 20)
            )

            chart = ReportService._create_kpi_chart(
                summary
            )

            if chart is not None:
                story.append(
                    Image(
                        chart,
                        width=480,
                        height=270,
                    )
                )

                story.append(
                    Spacer(1, 15)
                )

        # -----------------------------------------------------
        # RISK
        # -----------------------------------------------------

        if request.include_risk:

            story.append(
                Paragraph(
                    "Risk Intelligence",
                    styles["Heading2"],
                )
            )

            risk_level = summary.get(
                "overall_risk_level",
                "N/A",
            )

            risk_score = summary.get(
                "overall_risk_score",
                0,
            )

            story.append(
                Paragraph(
                    f"Risk Level: "
                    f"<b>{str(risk_level).upper()}</b>",
                    styles["Normal"],
                )
            )

            if risk_level == "N/A":
                story.append(
                    Paragraph(
                        "Risk Score: N/A — insufficient "
                        "operational data.",
                        styles["Normal"],
                    )
                )
            else:
                story.append(
                    Paragraph(
                        f"Risk Score: "
                        f"{float(risk_score):.2f}",
                        styles["Normal"],
                    )
                )

            drivers = summary.get(
                "primary_risk_drivers",
                [],
            )

            story.append(
                Paragraph(
                    "Primary Drivers: "
                    + (
                        ", ".join(
                            map(str, drivers)
                        )
                        if drivers
                        else "N/A"
                    ),
                    styles["Normal"],
                )
            )

            story.append(
                Spacer(1, 12)
            )

        # -----------------------------------------------------
        # ALERTS
        # -----------------------------------------------------

        if request.include_alerts:

            story.append(
                Paragraph(
                    "Alerts",
                    styles["Heading2"],
                )
            )

            alerts = summary.get(
                "alerts",
                [],
            )

            if alerts:

                for alert in alerts:
                    story.append(
                        Paragraph(
                            str(alert),
                            styles["Normal"],
                        )
                    )

            else:

                story.append(
                    Paragraph(
                        "No alerts available.",
                        styles["Normal"],
                    )
                )

            story.append(
                Spacer(1, 12)
            )

        # -----------------------------------------------------
        # RECOMMENDATIONS
        # -----------------------------------------------------

        if request.include_recommendations:

            story.append(
                Paragraph(
                    "Recommendations",
                    styles["Heading2"],
                )
            )

            recommendations = summary.get(
                "recommendations",
                [],
            )

            if recommendations:

                for recommendation in recommendations:
                    story.append(
                        Paragraph(
                            str(recommendation),
                            styles["Normal"],
                        )
                    )

            else:

                story.append(
                    Paragraph(
                        "No recommendations available.",
                        styles["Normal"],
                    )
                )

        document.build(story)

        buffer.seek(0)

        return buffer

    # =========================================================
    # DOCX
    # =========================================================

    @staticmethod
    def generate_docx(
        db,
        organization_id,
        request,
    ) -> BytesIO:

        summary = ReportService._summary(
            db,
            organization_id,
        )

        document = Document()

        document.add_heading(
            request.title,
            0,
        )

        document.add_paragraph(
            "Generated by ChainPulse AI"
        )

        document.add_paragraph(
            datetime.now().strftime(
                "Generated: %d %B %Y, %H:%M"
            )
        )

        # -----------------------------------------------------
        # DASHBOARD
        # -----------------------------------------------------

        if request.include_dashboard:

            document.add_heading(
                "Dashboard Overview",
                level=1,
            )

            metrics = ReportService._metrics(
                summary
            )

            table = document.add_table(
                rows=1,
                cols=2,
            )

            table.style = "Table Grid"

            table.rows[0].cells[0].text = "Metric"
            table.rows[0].cells[1].text = "Value"

            for label, value in metrics:

                cells = table.add_row().cells

                cells[0].text = str(label)
                cells[1].text = str(value)

            chart = ReportService._create_kpi_chart(
                summary
            )

            if chart is not None:

                document.add_paragraph()

                document.add_picture(
                    chart,
                    width=Inches(6.5),
                )

        # -----------------------------------------------------
        # RISK
        # -----------------------------------------------------

        if request.include_risk:

            document.add_heading(
                "Risk Intelligence",
                level=1,
            )

            risk_level = summary.get(
                "overall_risk_level",
                "N/A",
            )

            risk_score = summary.get(
                "overall_risk_score",
                0,
            )

            document.add_paragraph(
                f"Risk Level: "
                f"{str(risk_level).upper()}"
            )

            if risk_level == "N/A":

                document.add_paragraph(
                    "Risk Score: N/A — insufficient "
                    "operational data."
                )

            else:

                document.add_paragraph(
                    f"Risk Score: "
                    f"{float(risk_score):.2f}"
                )

            drivers = summary.get(
                "primary_risk_drivers",
                [],
            )

            document.add_paragraph(
                "Primary Drivers: "
                + (
                    ", ".join(
                        map(str, drivers)
                    )
                    if drivers
                    else "N/A"
                )
            )

        # -----------------------------------------------------
        # ALERTS
        # -----------------------------------------------------

        if request.include_alerts:

            document.add_heading(
                "Alerts",
                level=1,
            )

            alerts = summary.get(
                "alerts",
                [],
            )

            if alerts:

                for alert in alerts:
                    document.add_paragraph(
                        str(alert),
                        style="List Bullet",
                    )

            else:

                document.add_paragraph(
                    "No alerts available."
                )

        # -----------------------------------------------------
        # RECOMMENDATIONS
        # -----------------------------------------------------

        if request.include_recommendations:

            document.add_heading(
                "Recommendations",
                level=1,
            )

            recommendations = summary.get(
                "recommendations",
                [],
            )

            if recommendations:

                for recommendation in recommendations:
                    document.add_paragraph(
                        str(recommendation),
                        style="List Bullet",
                    )

            else:

                document.add_paragraph(
                    "No recommendations available."
                )

        buffer = BytesIO()

        document.save(buffer)

        buffer.seek(0)

        return buffer

    # =========================================================
    # EXCEL
    # =========================================================

    @staticmethod
    def generate_excel(
        db,
        organization_id,
        request,
    ) -> BytesIO:

        summary = ReportService._summary(
            db,
            organization_id,
        )

        workbook = Workbook()

        # -----------------------------------------------------
        # SUMMARY SHEET
        # -----------------------------------------------------

        sheet = workbook.active
        sheet.title = "Dashboard"

        sheet["A1"] = request.title
        sheet["A1"].font = Font(
            bold=True,
            size=18,
        )

        sheet["A2"] = "Generated by ChainPulse AI"
        sheet["A3"] = datetime.now().strftime(
            "%d %B %Y, %H:%M"
        )

        sheet["A5"] = "Metric"
        sheet["B5"] = "Value"

        header_fill = PatternFill(
            fill_type="solid",
            fgColor="1F2937",
        )

        for cell in sheet[5]:

            cell.font = Font(
                bold=True,
                color="FFFFFF",
            )

            cell.fill = header_fill

            cell.alignment = Alignment(
                horizontal="center"
            )

        metrics = ReportService._metrics(
            summary
        )

        row = 6

        for label, value in metrics:

            sheet.cell(
                row=row,
                column=1,
                value=label,
            )

            sheet.cell(
                row=row,
                column=2,
                value=value,
            )

            row += 1

        sheet.column_dimensions["A"].width = 32
        sheet.column_dimensions["B"].width = 24

        # -----------------------------------------------------
        # OPERATIONAL CHART
        # -----------------------------------------------------

        if ReportService._has_operational_data(
            summary
        ):

            chart_sheet = workbook.create_sheet(
                "Charts"
            )

            chart_sheet["A1"] = (
                "Operational Overview"
            )

            chart_sheet["A1"].font = Font(
                bold=True,
                size=16,
            )

            chart_sheet["A3"] = "Category"
            chart_sheet["B3"] = "Count"

            operational_data = [
                (
                    "Products",
                    summary.get(
                        "product_count",
                        0,
                    ),
                ),
                (
                    "Demand Records",
                    summary.get(
                        "demand_record_count",
                        0,
                    ),
                ),
                (
                    "Inventory Records",
                    summary.get(
                        "inventory_record_count",
                        0,
                    ),
                ),
                (
                    "Suppliers",
                    summary.get(
                        "supplier_count",
                        0,
                    ),
                ),
            ]

            for index, (
                label,
                value,
            ) in enumerate(
                operational_data,
                start=4,
            ):

                chart_sheet.cell(
                    row=index,
                    column=1,
                    value=label,
                )

                chart_sheet.cell(
                    row=index,
                    column=2,
                    value=value,
                )

            chart = BarChart()

            chart.title = (
                "Operational Overview"
            )

            chart.y_axis.title = "Count"
            chart.x_axis.title = "Category"

            data = Reference(
                chart_sheet,
                min_col=2,
                min_row=3,
                max_row=3 + len(
                    operational_data
                ),
            )

            categories = Reference(
                chart_sheet,
                min_col=1,
                min_row=4,
                max_row=3 + len(
                    operational_data
                ),
            )

            chart.add_data(
                data,
                titles_from_data=True,
            )

            chart.set_categories(
                categories
            )

            chart.height = 8
            chart.width = 14

            chart_sheet.add_chart(
                chart,
                "D3",
            )

        else:

            empty_sheet = workbook.create_sheet(
                "Charts"
            )

            empty_sheet["A1"] = (
                "No operational data available."
            )

            empty_sheet["A1"].font = Font(
                italic=True,
            )

        # -----------------------------------------------------
        # RISK SHEET
        # -----------------------------------------------------

        if request.include_risk:

            risk_sheet = workbook.create_sheet(
                "Risk"
            )

            risk_sheet["A1"] = (
                "Risk Intelligence"
            )

            risk_sheet["A1"].font = Font(
                bold=True,
                size=16,
            )

            risk_sheet["A3"] = "Metric"
            risk_sheet["B3"] = "Value"

            risk_values = [
                (
                    "Overall Risk",
                    summary.get(
                        "overall_risk_level",
                        "N/A",
                    ),
                ),
                (
                    "Risk Score",
                    (
                        summary.get(
                            "overall_risk_score",
                            0,
                        )
                        if summary.get(
                            "overall_risk_level"
                        ) != "N/A"
                        else "N/A"
                    ),
                ),
                (
                    "Demand Risk",
                    summary.get(
                        "demand_risk_level",
                        "N/A",
                    ),
                ),
                (
                    "Inventory Risk",
                    summary.get(
                        "inventory_risk_level",
                        "N/A",
                    ),
                ),
                (
                    "Supplier Risk",
                    summary.get(
                        "supplier_risk_level",
                        "N/A",
                    ),
                ),
            ]

            for index, (
                label,
                value,
            ) in enumerate(
                risk_values,
                start=4,
            ):

                risk_sheet.cell(
                    index,
                    1,
                    label,
                )

                risk_sheet.cell(
                    index,
                    2,
                    value,
                )

            risk_sheet.column_dimensions[
                "A"
            ].width = 28

            risk_sheet.column_dimensions[
                "B"
            ].width = 24

        # -----------------------------------------------------
        # ALERTS
        # -----------------------------------------------------

        if request.include_alerts:

            alert_sheet = workbook.create_sheet(
                "Alerts"
            )

            alert_sheet["A1"] = "Alerts"
            alert_sheet["A1"].font = Font(
                bold=True,
                size=16,
            )

            alerts = summary.get(
                "alerts",
                [],
            )

            if alerts:

                for index, alert in enumerate(
                    alerts,
                    start=3,
                ):

                    alert_sheet.cell(
                        index,
                        1,
                        str(alert),
                    )

            else:

                alert_sheet["A3"] = (
                    "No alerts available."
                )

            alert_sheet.column_dimensions[
                "A"
            ].width = 100

        # -----------------------------------------------------
        # RECOMMENDATIONS
        # -----------------------------------------------------

        if request.include_recommendations:

            recommendation_sheet = (
                workbook.create_sheet(
                    "Recommendations"
                )
            )

            recommendation_sheet["A1"] = (
                "Recommendations"
            )

            recommendation_sheet["A1"].font = Font(
                bold=True,
                size=16,
            )

            recommendations = summary.get(
                "recommendations",
                [],
            )

            if recommendations:

                for index, recommendation in enumerate(
                    recommendations,
                    start=3,
                ):

                    recommendation_sheet.cell(
                        index,
                        1,
                        str(recommendation),
                    )

            else:

                recommendation_sheet["A3"] = (
                    "No recommendations available."
                )

            recommendation_sheet.column_dimensions[
                "A"
            ].width = 100

        # -----------------------------------------------------
        # SAVE
        # -----------------------------------------------------

        buffer = BytesIO()

        workbook.save(buffer)

        buffer.seek(0)

        return buffer